#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AGENTE DE DISEÑO DE CONDENSADORES DE CASCO Y TUBOS
Transferencia de Energía Térmica - Unidad 3 (Condensación)

Qué hace (3 fases del TP):
  FASE I   Inputs: se cargan en un diccionario (CASO_EJEMPLO) o en un .json
  FASE II  Algoritmo iterativo: balance, DMLT, Ft (suma carcasas), prediseño,
           ho / hi, Uc, Ud, Rd, pérdidas de carga. Si algo no cumple, el
           agente cambia solo: carcasas, diámetro de casco, pasos, espaciado de
           deflectores y longitud de tubo, y vuelve a evaluar.
  FASE III Salida: Hoja de Especificaciones TEMA (.xlsx) + log (.txt) + .json

Uso:
  python agente_condensador.py                 -> corre el caso de ejemplo
  python agente_condensador.py --plantilla     -> escribe caso_plantilla.json
  python agente_condensador.py mi_caso.json    -> corre tu caso
  python agente_condensador.py mi_caso.json -o salida

Unidades: SI (kg/s, °C, Pa, W, J/kg, m). Los inputs de caudal se dan en kg/s.
"""
import math
import json
import argparse
import os
import copy
from datetime import date

g = 9.80665

# =============================================================================
# FASE I - DATOS DE ENTRADA (CASO DE EJEMPLO / VALIDACIÓN)
# Vapor de agua saturado a 3 bar abs condensando en el casco, agua por tubos.
# Propiedades del vapor y del condensado tomadas de tablas de vapor.
# =============================================================================
CASO_EJEMPLO = {
    "identificacion": {
        "cliente": "Cátedra Transferencia de Energía Térmica",
        "item": "E-101",
        "servicio": "Condensador de vapor de agua",
        "tipo_TEMA": "BEM",          # BEM, AEL, AES, AET (fijo/flotante) o BEU/AEU (tubos en U)
        "orientacion": "Horizontal, condensación en casco",
    },
    "caliente": {                    # corriente que condensa (lado casco)
        "fluido": "Vapor de agua",
        "W": 0.20,                   # caudal másico de vapor [kg/s]
        "T1": 133.5, "T2": 133.5,    # T entrada / salida [°C]. T1>Tsat: sobrecalentado; T2<Tsat: subenfriado
        "Tsat": 133.5,               # temperatura de saturación a la presión de operación [°C]
        "P_oper": 300e3,             # presión de operación [Pa abs]
        "lam": 2163.5e3,             # calor latente [J/kg]
        "cp_v": 2170.0, "rho_v": 1.651, "mu_v": 1.34e-5, "k_v": 0.0275,       # vapor
        "cp_l": 4250.0, "rho_l": 943.0, "mu_l": 2.32e-4, "k_l": 0.685,        # condensado
        "T_film_ref": 120.0,         # T a la que se dieron las props. del condensado [°C]
        "Rd": 0.0000876,             # ensuciamiento requerido lado casco [m2K/W]
        "dP_adm": 15e3,              # pérdida de carga admisible lado casco [Pa]
        "P_diseno": 600e3, "T_diseno": 170.0,
    },
    "fria": {                        # refrigerante (lado tubos)
        "fluido": "Agua de enfriamiento",
        "t1": 25.0, "t2": 45.0,      # t2 = None -> se calcula a partir de W
        "W": None,                   # caudal [kg/s]; None -> se calcula a partir de t2
        "props": "agua",             # "agua" = correlaciones internas; o "manual" con props_man
        "props_man": {"cp": 4182.0, "rho": 995.0, "mu": 7.2e-4, "k": 0.62},
        "Rd": 0.000176,              # ensuciamiento lado tubos [m2K/W]
        "dP_adm": 70e3,              # pérdida de carga admisible lado tubos [Pa]
        "P_oper": 300e3, "P_diseno": 600e3, "T_diseno": 70.0,
    },
    "constructivo": {
        "Do": 0.01905, "Di": 0.01575,     # 3/4" BWG 16
        "L_max": 3.66,                    # longitud máxima de tubo [m] (12 ft)
        "arreglo": "triangular",          # "triangular" (60°) o "cuadrado" (90°)
        "Pt": 0.0238125,                  # paso entre tubos [m]
        "k_pared": 50.0,                  # conductividad del tubo [W/mK] (acero al carbono)
        "material_tubo": "Acero al carbono SA-179",
        "material_casco": "Acero al carbono SA-516 Gr70",
        "corte_deflector": 0.25,
        "v_min": 1.0, "v_max": 3.0,       # velocidad admisible por tubos [m/s]
    },
    "diseno": {
        "Ud_est": 1500.0,            # coeficiente global estimado [W/m2K] (150-300 BTU/h ft2 F ~ 850-1700)
        "Ft_min": 0.75,
        "u_corte_mueller": 20.0,     # velocidad de vapor [m/s] desde la que se evalúa corte (Mueller)
        "mueller_factor_max": 3.0,   # tope razonable h_Mueller/h_f
    },
}

# =============================================================================
# TABLAS Y CONSTANTES
# =============================================================================
# Diámetros internos normalizados de carcasa [m] (4" ... 39")
DS_TEMA = [0.1023, 0.1541, 0.2032, 0.2540, 0.3048, 0.3366, 0.3874, 0.4382, 0.4890, 0.5398, 0.5906,
           0.6350, 0.6858, 0.7366, 0.7874, 0.8382, 0.8890, 0.9398, 0.9906]

# Espesores de pared por BWG [mm]
BWG = {10: 3.40, 11: 3.05, 12: 2.77, 13: 2.41, 14: 2.11, 15: 1.83, 16: 1.65,
       17: 1.47, 18: 1.24, 19: 1.07, 20: 0.89}

# Constantes de conteo de tubos Nt = K1 (Db/Do)^n1 para Pt = 1.25 Do  (Coulson & Richardson, tabla 12.4)
K_TUBOS = {("triangular", 1): (0.319, 2.142), ("triangular", 2): (0.249, 2.207),
           ("triangular", 4): (0.175, 2.285), ("triangular", 6): (0.0743, 2.499),
           ("triangular", 8): (0.0365, 2.675),
           ("cuadrado", 1): (0.215, 2.207), ("cuadrado", 2): (0.156, 2.291),
           ("cuadrado", 4): (0.158, 2.263), ("cuadrado", 6): (0.0402, 2.617),
           ("cuadrado", 8): (0.0331, 2.643)}

# Diámetros internos de boquillas [m] (DN 50 ... DN 600)
BOQUILLAS = [("2\"", 0.0525), ("2.5\"", 0.0627), ("3\"", 0.0779), ("4\"", 0.1023),
             ("5\"", 0.1282), ("6\"", 0.1541), ("8\"", 0.2027), ("10\"", 0.2545),
             ("12\"", 0.3048), ("14\"", 0.3366), ("16\"", 0.3874), ("18\"", 0.4382),
             ("20\"", 0.4890), ("24\"", 0.5900)]


# =============================================================================
# PROPIEDADES DEL AGUA LÍQUIDA (refrigerante) - correlaciones, T en °C
# =============================================================================
def props_agua(T):
    Tk = T + 273.15
    mu = 2.414e-5 * 10 ** (247.8 / (Tk - 140.0))
    k = 0.5706 + 0.001756 * T - 6.46e-6 * T ** 2
    rho = 1000.0 * (1 - (T + 288.9414) / (508929.2 * (T + 68.12963)) * (T - 3.9863) ** 2)
    return {"cp": 4182.0, "rho": rho, "mu": mu, "k": k}


def props_fria(f, T):
    if f.get("props", "agua") == "agua":
        return props_agua(T)
    return f["props_man"]


# =============================================================================
# FUNCIONES TÉRMICAS BÁSICAS
# =============================================================================
def lmtd(a, b):
    """DMLT a partir de las dos diferencias terminales (contracorriente)."""
    if a <= 0 or b <= 0:
        raise ValueError("Cruce de temperaturas: diferencia terminal <= 0. Revisar T de proceso.")
    if abs(a - b) < 1e-9:
        return a
    return (a - b) / math.log(a / b)


def ft_factor(R, S, N):
    """Factor Ft para N carcasas en serie (1-2 cada una). R=(T1-T2)/(t2-t1), S=(t2-t1)/(T1-t1)."""
    if R < 1e-9:
        return 1.0                      # un fluido isotérmico: Ft = 1
    try:
        if abs(R - 1) < 1e-6:
            P1 = S / (N - (N - 1) * S)
            if not 0 < P1 < 1:
                return 0.0
            s2 = math.sqrt(2.0)
            num = s2 * P1 / (N * (1 - P1))
            den = math.log((2 - P1 * (2 - s2)) / (2 - P1 * (2 + s2)))
            return num / den
        z = ((1 - S * R) / (1 - S)) ** (1.0 / N)
        P1 = (1 - z) / (R - z)
        if not 0 < P1 < 1:
            return 0.0
        r2 = math.sqrt(R * R + 1)
        num = r2 / (N * (R - 1)) * math.log((1 - P1) / (1 - R * P1))
        a = (2 / P1 - 1 - R + r2) / (2 / P1 - 1 - R - r2)
        if a <= 0:
            return 0.0
        return max(0.0, num / math.log(a))
    except (ValueError, ZeroDivisionError):
        return 0.0


def balance(c, log):
    """Balance de energía (calor sensible + latente), caudal/temperatura de agua y zonas."""
    h, f = c["caliente"], c["fria"]
    Wh, Ts, T1, T2 = h["W"], h["Tsat"], h["T1"], h["T2"]
    Qd = Wh * h["cp_v"] * (T1 - Ts) if T1 > Ts + 1e-6 else 0.0      # desobrecalentamiento
    Ql = Wh * h["lam"]                                              # calor latente
    Qs = Wh * h["cp_l"] * (Ts - T2) if T2 < Ts - 1e-6 else 0.0      # subenfriamiento
    Q = Qd + Ql + Qs
    t1 = f["t1"]
    if f.get("t2") is not None:
        t2 = f["t2"]
        cp = props_fria(f, (t1 + t2) / 2)["cp"]
        Wc = Q / (cp * (t2 - t1))
    else:
        Wc = f["W"]
        t2 = t1 + 10
        for _ in range(5):
            cp = props_fria(f, (t1 + t2) / 2)["cp"]
            t2 = t1 + Q / (Wc * cp)
    cp = props_fria(f, (t1 + t2) / 2)["cp"]
    tb1 = t1 + Qs / (Wc * cp)                      # T del agua al salir de la zona de subenfriamiento
    zonas = []
    if Qs > 0:
        zonas.append({"nombre": "Subenfriamiento", "Q": Qs, "hot_in": Ts, "hot_out": T2,
                      "c_in": t1, "c_out": tb1, "dTml": lmtd(Ts - tb1, T2 - t1)})
    zonas.append({"nombre": "Condensación" + (" + desobrecalentamiento" if Qd > 0 else ""),
                  "Q": Ql + Qd, "hot_in": T1, "hot_out": Ts, "c_in": tb1, "c_out": t2,
                  "dTml": lmtd(Ts - tb1, Ts - t2)})
    log.append(f"Balance: Q = {Q/1e3:,.1f} kW (latente {Ql/1e3:,.1f}, desobrecal. {Qd/1e3:,.1f}, "
               f"subenfr. {Qs/1e3:,.1f}); Wc = {Wc:.3f} kg/s; t2 = {t2:.1f} °C")
    for z in zonas:
        log.append(f"  Zona {z['nombre']}: Q = {z['Q']/1e3:,.1f} kW, DMLT = {z['dTml']:.2f} K")
    return {"Q": Q, "Qd": Qd, "Ql": Ql, "Qs": Qs, "Wc": Wc, "t1": t1, "t2": t2,
            "tm": (t1 + t2) / 2, "zonas": zonas}


def ft_global(bal, c, Ns, n):
    """Ft de cada zona para Ns carcasas y n pasos por tubo (n=1 -> contracorriente pura)."""
    fts = []
    for z in bal["zonas"]:
        if n == 1 or z["hot_in"] == z["hot_out"]:
            fts.append(1.0)
            continue
        R = (z["hot_in"] - z["hot_out"]) / (z["c_out"] - z["c_in"])
        S = (z["c_out"] - z["c_in"]) / (z["hot_in"] - z["c_in"])
        fts.append(ft_factor(R, S, Ns))
    return fts


# =============================================================================
# COEFICIENTES PELICULARES Y PÉRDIDAS DE CARGA
# =============================================================================
def hi_tubos(Re, Pr, k, Di, Ltot):
    """Coeficiente interno: laminar (Hausen), transición (Gnielinski), turbulento (Dittus-Boelter)."""
    if Re < 2300:
        Gz = Di / Ltot * Re * Pr
        Nu = 3.66 + 0.0668 * Gz / (1 + 0.04 * Gz ** (2 / 3))
    elif Re < 10000:
        f = (0.79 * math.log(Re) - 1.64) ** -2
        Nu = (f / 8) * (Re - 1000) * Pr / (1 + 12.7 * math.sqrt(f / 8) * (Pr ** (2 / 3) - 1))
    else:
        Nu = 0.023 * Re ** 0.8 * Pr ** 0.4
    return Nu * k / Di


def h_condensacion(h, Nt, L, Gv, Do, Tw, dis, avisos):
    """
    Condensación en el exterior de un banco de tubos horizontales.
      Re_pelicula = 4 G'' / mu,  G'' = W / (L Nt^(2/3))   (Kern, salpicaduras entre filas)
      Re <= 2100 -> McAdams horizontal:  h = 1.51 (k^3 rho (rho-rho_v) g / mu^2)^(1/3) Re^(-1/3)
      Re  > 2100 -> película turbulenta (forma de Colburn, equivalente a Dukler):  0.0077 Re^0.4
      Si la velocidad del vapor es alta -> término convectivo de Mueller (Cao, pág. 257).
    """
    kl, rl, rv, mul = h["k_l"], h["rho_l"], h["rho_v"], h["mu_l"]
    base = (kl ** 3 * rl * (rl - rv) * g / mul ** 2) ** (1 / 3)
    Gpp = h["W"] / (L * Nt ** (2 / 3))
    Re_f = 4 * Gpp / mul
    if Re_f <= 2100:
        hf, modelo = 1.51 * base * Re_f ** (-1 / 3), "McAdams (película laminar, tubos horizontales)"
    else:
        hf, modelo = 0.0077 * base * Re_f ** 0.4, "Película turbulenta (Colburn/Dukler)"
    u_v = Gv / rv
    hfinal = hf
    if u_v >= dis["u_corte_mueller"]:
        R = rl * mul / (rv * h["mu_v"])
        Pr = h["cp_l"] * mul / kl
        H = h["cp_l"] * max(h["Tsat"] - Tw, 0.1) / (Pr * h["lam"])
        X = 0.9 * (1 + 1 / (R * H)) ** (1 / 3)
        # ATENCIÓN: Re_LG tal como figura en la teoría (Cao p.257). Verificar contra el libro.
        Re_LG = Do * Gv * rl / (mul * rv)
        Nu_f = hf * Do / kl
        hm = ((X ** 4 * Re_LG ** 2 + Nu_f ** 4) ** 0.25) * kl / Do
        if hm > dis["mueller_factor_max"] * hf:
            avisos.append(f"Mueller dio h={hm:,.0f} (> {dis['mueller_factor_max']:.0f} x h_f={hf:,.0f}): "
                          "resultado dudoso, se adopta h_f (conservador). Verificar Re_LG en Cao p.257.")
        else:
            hfinal, modelo = hm, "Mueller (corte de vapor) sobre Dukler"
    return hfinal, hf, Re_f, modelo, u_v


def ho_casco_monofasico(h, Gs, De, Ts_mean=None):
    """Zona de subenfriamiento (líquido en casco): Kern  Nu = 0.36 Re^0.55 Pr^(1/3)."""
    Re = De * Gs / h["mu_l"]
    Pr = h["cp_l"] * h["mu_l"] / h["k_l"]
    return 0.36 * h["k_l"] / De * Re ** 0.55 * Pr ** (1 / 3)


def elegir_boquilla(m, rho, criterio, limite):
    for nom, d in BOQUILLAS:
        v = m / (rho * math.pi / 4 * d * d)
        ok = (rho * v * v <= limite) if criterio == "rhov2" else (v <= limite)
        if ok:
            return nom, d, v
    nom, d = BOQUILLAS[-1]
    return nom, d, m / (rho * math.pi / 4 * d * d)


def holgura_casco(tipo):
    t = tipo.upper()
    if t.endswith("U"):
        return 0.020
    if t.endswith("S"):
        return 0.055
    if t.endswith("T"):
        return 0.090
    return 0.013            # cabezal fijo (M, L, N)


def n_tubos(Ds, n, arreglo, Do, Pt, tipo):
    if tipo.upper().endswith("U") and n % 2:
        return 0
    K1, n1 = K_TUBOS[(arreglo, n)]
    Db = Ds - holgura_casco(tipo)
    if Db <= 2 * Do:
        return 0
    return int(K1 * (Db / Do) ** n1 * (1.25 * Do / Pt) ** 2)


def bwg_de(Do, Di):
    pared = (Do - Di) / 2 * 1000
    return min(BWG, key=lambda b: abs(BWG[b] - pared))


# =============================================================================
# EVALUACIÓN DE UNA CONFIGURACIÓN GEOMÉTRICA COMPLETA
# =============================================================================
def evaluar(c, bal, Ns, Ds, n, Bf, L):
    h, f, k, dis = c["caliente"], c["fria"], c["constructivo"], c["diseno"]
    tipo = c["identificacion"]["tipo_TEMA"]
    Do, Di, Pt, kw = k["Do"], k["Di"], k["Pt"], k["k_pared"]
    Nt = n_tubos(Ds, n, k["arreglo"], Do, Pt, tipo)
    if Nt < 2:
        return None
    avisos = []
    A_disp = Nt * math.pi * Do * L
    Wh, Wc = h["W"], bal["Wc"]

    # ---------------- lado tubos ----------------
    pw = props_fria(f, bal["tm"])
    a_t = Nt / n * math.pi * Di ** 2 / 4
    v = Wc / (pw["rho"] * a_t)
    Re_t = pw["rho"] * v * Di / pw["mu"]
    Pr_t = pw["cp"] * pw["mu"] / pw["k"]
    hi = hi_tubos(Re_t, Pr_t, pw["k"], Di, L * n)
    hio = hi * Di / Do
    fd = 64 / Re_t if Re_t < 2300 else (0.79 * math.log(Re_t) - 1.64) ** -2
    q_din = pw["rho"] * v * v / 2
    dP_fric = fd * (L * n / Di) * q_din
    dP_ret = 4 * n * q_din                             # cambio de dirección en cabezales (Kern)
    nom_tn, d_tn, v_tn = elegir_boquilla(Wc, pw["rho"], "v", 3.0)
    dP_noz_t = 1.5 * pw["rho"] * v_tn ** 2 / 2         # 1.0 entrada + 0.5 salida
    dP_t = Ns * (dP_fric + dP_ret + dP_noz_t)

    # ---------------- lado casco: geometría ----------------
    Cp = Pt - Do
    nb = max(1, math.ceil(L / (Bf * Ds)))
    B = L / nb
    As = Ds * Cp * B / Pt
    if k["arreglo"] == "cuadrado":
        De = 4 * (Pt ** 2 - math.pi * Do ** 2 / 4) / (math.pi * Do)
    else:
        De = 4 * (0.5 * Pt * 0.866 * Pt - 0.5 * math.pi * Do ** 2 / 4) / (0.5 * math.pi * Do)
    Gs = Wh / As

    # ---------------- coeficientes por zona y áreas ----------------
    def U_limpio(ho_):
        return 1.0 / (1.0 / ho_ + Do * math.log(Do / Di) / (2 * kw) + 1.0 / hio)

    fts = ft_global(bal, c, Ns, n)
    zonas = []
    Tw = h["Tsat"] - 10.0
    ho_c, hf, Re_f, modelo, u_v = h_condensacion(h, Nt, L, Gs, Do, Tw, dis, avisos)
    for z, ft in zip(bal["zonas"], fts):
        if z["nombre"].startswith("Subenf"):
            ho_z = ho_casco_monofasico(h, Gs, De)
        else:
            ho_z = ho_c
        Uc_z = U_limpio(ho_z)
        A_z = z["Q"] / (Uc_z * z["dTml"] * max(ft, 1e-6))
        zonas.append({**z, "ho": ho_z, "Uc": Uc_z, "Ft": ft, "A_limpia": A_z})
    A_limp = sum(z["A_limpia"] for z in zonas)
    Uc = sum(z["Uc"] * z["A_limpia"] for z in zonas) / A_limp
    dT_bal = bal["Q"] / sum(z["Q"] / (z["dTml"] * max(z["Ft"], 1e-6)) for z in zonas)
    Ud = bal["Q"] / (A_disp * dT_bal)

    # temperatura de pared (afina Mueller y chequea T de película)
    zc = zonas[-1]
    q_c = zc["Q"] / (A_disp * zc["A_limpia"] / A_limp)
    Tw = h["Tsat"] - q_c / ho_c
    if u_v >= dis["u_corte_mueller"]:
        avisos.clear()
        ho_c, hf, Re_f, modelo, u_v = h_condensacion(h, Nt, L, Gs, Do, Tw, dis, avisos)
        for zz in zonas:
            if not zz["nombre"].startswith("Subenf"):
                zz["ho"] = ho_c
                zz["Uc"] = U_limpio(ho_c)
        # se recalculan áreas con el nuevo ho
        for zz in zonas:
            zz["A_limpia"] = zz["Q"] / (zz["Uc"] * zz["dTml"] * max(zz["Ft"], 1e-6))
        A_limp = sum(zz["A_limpia"] for zz in zonas)
        Uc = sum(zz["Uc"] * zz["A_limpia"] for zz in zonas) / A_limp
        dT_bal = bal["Q"] / sum(zz["Q"] / (zz["dTml"] * max(zz["Ft"], 1e-6)) for zz in zonas)
        Ud = bal["Q"] / (A_disp * dT_bal)
    T_film = (h["Tsat"] + Tw) / 2
    if abs(T_film - h["T_film_ref"]) > 10:
        avisos.append(f"T de película calculada {T_film:.0f} °C difiere de la de las propiedades "
                      f"({h['T_film_ref']:.0f} °C): reevaluar propiedades del condensado.")

    Rd_calc = (Uc - Ud) / (Uc * Ud)
    Rd_req = h["Rd"] + f["Rd"] * Do / Di

    # ---------------- pérdida de carga lado casco ----------------
    Re_s = De * Gs / h["mu_v"]
    fs = math.exp(0.576 - 0.19 * math.log(Re_s))
    dP_haz = 0.5 * fs * Gs ** 2 * Ds * (L / B) / (2 * h["rho_v"] * De)   # 0.5: vapor que condensa
    nom_sv, d_sv, v_sv = elegir_boquilla(Wh, h["rho_v"], "rhov2", 2232.0)  # criterio TEMA rho v^2
    nom_sc, d_sc, v_sc = elegir_boquilla(Wh, h["rho_l"], "v", 1.0)
    dP_noz_s = 1.0 * h["rho_v"] * v_sv ** 2 / 2
    dP_s = dP_haz + dP_noz_s

    ft_min = min(fts)
    viol = {
        "Ft": max(0.0, (dis["Ft_min"] - ft_min) / dis["Ft_min"]),
        "Rd": max(0.0, (Rd_req - Rd_calc) / Rd_req) if Rd_calc < Rd_req else 0.0,
        "v_tubos": max(0.0, (k["v_min"] - v) / k["v_min"], (v - k["v_max"]) / k["v_max"]),
        "dP_tubos": max(0.0, (dP_t - f["dP_adm"]) / f["dP_adm"]),
        "dP_casco": max(0.0, (dP_s - h["dP_adm"]) / h["dP_adm"]),
    }
    return {
        "Ns": Ns, "Ds": Ds, "n": n, "Nt": Nt, "L": L, "B": B, "n_deflectores": nb - 1,
        "A_disp": A_disp, "A_total": A_disp * Ns, "A_limpia": A_limp,
        "exceso": (1.0 / (1.0 / Uc + Rd_req) / Ud - 1) * 100,   # A_disp / A_req(sucia) - 1
        "hi": hi, "hio": hio, "ho": ho_c, "ho_f": hf, "modelo_ho": modelo, "Re_pelicula": Re_f,
        "Uc": Uc, "Ud": Ud, "Rd_calc": Rd_calc, "Rd_req": Rd_req, "dT_bal": dT_bal,
        "Ft": ft_min, "zonas": zonas,
        "v_tubos": v, "Re_tubos": Re_t, "Pr_tubos": Pr_t, "Tw": Tw, "T_film": T_film,
        "dP_tubos": dP_t, "dP_tubos_fric": dP_fric * Ns, "dP_tubos_ret": dP_ret * Ns,
        "dP_tubos_noz": dP_noz_t * Ns,
        "dP_casco": dP_s, "dP_casco_haz": dP_haz, "dP_casco_noz": dP_noz_s,
        "u_vapor_casco": u_v, "Gs": Gs, "De": De, "As": As, "Re_casco": Re_s,
        "boq_tubos": (nom_tn, d_tn, v_tn), "boq_vapor": (nom_sv, d_sv, v_sv),
        "boq_condensado": (nom_sc, d_sc, v_sc), "viol": viol,
        "cumple": all(x <= 1e-9 for x in viol.values()), "avisos": avisos,
        "propiedades_agua": pw,
    }


# =============================================================================
# FASE II - NÚCLEO ITERATIVO
# =============================================================================
def disenar(c, log):
    dis, k = c["diseno"], c["constructivo"]
    bal = balance(c, log)

    # 1) Ft: si < Ft_min con 2 pasos, el agente suma carcasas en serie
    Ns = 1
    while True:
        ft = min(ft_global(bal, c, Ns, 2))
        log.append(f"Chequeo Ft con {Ns} carcasa(s) en serie y 2 pasos: Ft = {ft:.3f}")
        if ft >= dis["Ft_min"] or Ns >= 6:
            break
        log.append(f"  Ft < {dis['Ft_min']} -> se agrega una carcasa en serie")
        Ns += 1

    # 2) prediseño con Ud estimado
    Ft_pre = max(ft, 0.01)
    dT_pre = bal["Q"] / sum(z["Q"] / z["dTml"] for z in bal["zonas"])
    A_prov = bal["Q"] / (dis["Ud_est"] * dT_pre * Ft_pre)
    Nt_prov = A_prov / (math.pi * k["Do"] * k["L_max"])
    log.append(f"Prediseño: Ud_est = {dis['Ud_est']:.0f} W/m2K, A' = {A_prov:.2f} m2, "
               f"Nt provisorio ~ {Nt_prov:.0f} tubos (L = {k['L_max']} m)")

    # 3) búsqueda: casco normalizado x pasos x espaciado de deflectores x longitud
    largos = [k["L_max"], 0.75 * k["L_max"], 0.5 * k["L_max"], 0.35 * k["L_max"]]
    pasos = [1, 2, 4, 6, 8]
    fracciones_B = [1.0, 0.75, 0.5, 0.4, 0.3, 0.2]
    factibles, mejor_inf = [], None
    log.append("Iteración sobre Ds / pasos / espaciado B / longitud (se reconfigura solo):")
    for Ds in DS_TEMA:
        causas, hay = {}, False
        for L in largos:
            for n in pasos:
                for Bf in fracciones_B:
                    r = evaluar(c, bal, Ns, Ds, n, Bf, L)
                    if r is None:
                        continue
                    if r["cumple"]:
                        factibles.append(r)
                        hay = True
                    else:
                        tot = sum(r["viol"].values())
                        if mejor_inf is None or tot < mejor_inf[0]:
                            mejor_inf = (tot, r)
                        for kk, vv in r["viol"].items():
                            if vv > 0:
                                causas[kk] = causas.get(kk, 0) + 1
        if hay:
            log.append(f"  Ds = {Ds*1000:.0f} mm -> hay configuraciones que cumplen todo")
        else:
            top = max(causas, key=causas.get) if causas else "sin haz posible"
            log.append(f"  Ds = {Ds*1000:.0f} mm -> descartado (causa dominante: {top})")
        if factibles:
            break           # el casco más chico que cumple es el óptimo en área
    if factibles:
        mejor = min(factibles, key=lambda r: (r["A_total"], r["dP_tubos"] + r["dP_casco"]))
        log.append(f"Selección: Ds = {mejor['Ds']*1000:.0f} mm, n = {mejor['n']}, Nt = {mejor['Nt']}, "
                   f"B = {mejor['B']*1000:.0f} mm, L = {mejor['L']:.2f} m")
    else:
        mejor = mejor_inf[1]
        mejor["avisos"].append("NINGUNA configuración cumple todas las restricciones: se muestra la de menor "
                               "violación. Revisar inputs (Δp admisibles, Rd, longitud máxima).")
        log.append("¡Sin solución factible! Se informa la configuración de menor violación.")
    return bal, mejor


def notas_criticas(c, r):
    """Observaciones de criterio de ingeniería generadas con los resultados."""
    h, k = c["caliente"], c["constructivo"]
    n = []
    if r["u_vapor_casco"] > 30:
        n.append(f"Velocidad de vapor en casco {r['u_vapor_casco']:.0f} m/s: riesgo de vibración inducida por flujo "
                 "y erosión de tubos; evaluar placas de impacto, deflectores de doble segmento o flujo dividido.")
    elif r["u_vapor_casco"] > 15:
        n.append(f"Velocidad de vapor en casco {r['u_vapor_casco']:.0f} m/s: elevada; verificar vibración.")
    if r["Re_pelicula"] > 2100:
        n.append("Película de condensado en régimen turbulento: se usó la correlación turbulenta.")
    n.append("Drenaje: condensado sale por el punto bajo del casco; ranura de drenaje en deflectores "
             "para que no se acumule líquido entre ellos y no se inunden tubos inferiores.")
    n.append("Venteo: prever venteo de no condensables en el extremo opuesto a la entrada de vapor "
             "(zona de mayor concentración); su acumulación baja ho y la T de condensación.")
    tipo = c["identificacion"]["tipo_TEMA"].upper()
    dT_met = h["Tsat"] - c["fria"]["t1"]
    if tipo.endswith("M") or tipo.endswith("L") or tipo.endswith("N"):
        if dT_met > 50:
            n.append(f"Cabezal fijo con diferencia de T casco/tubos hasta {dT_met:.0f} K: verificar tensiones "
                     "térmicas y evaluar junta de expansión o cabezal flotante / tubos en U.")
    n.append("Limpieza: agua por tubos permite limpieza mecánica del interior retirando el cabezal; "
             "el casco (vapor limpio) no requiere limpieza frecuente.")
    if r["exceso"] > 40:
        n.append(f"Exceso de área {r['exceso']:.0f}%: sobredimensionado; evaluar menor longitud o menor Ds.")
    if r["v_tubos"] > 2.5:
        n.append("Velocidad por tubos alta: riesgo de erosión-corrosión; considerar material más resistente.")
    n.append("Optimización económica: comparar con más caudal de agua (menor área, mayor bombeo) "
             "según costo de agua y de área.")
    return n + r["avisos"]


# =============================================================================
# FASE III - HOJA DE ESPECIFICACIONES TEMA (xlsx)
# =============================================================================
def escribir_hoja_tema(c, bal, r, notas, log, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    h, f, k, idn = c["caliente"], c["fria"], c["constructivo"], c["identificacion"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja TEMA"
    F = lambda **kw: Font(name="Arial", size=kw.pop("size", 9), **kw)
    gris = PatternFill("solid", fgColor="D9E1F2")
    azul = PatternFill("solid", fgColor="1F3864")
    borde = Border(*(Side(style="thin", color="808080"),) * 4)
    row = [1]

    def titulo(texto):
        ws.cell(row[0], 1, texto).font = F(bold=True, color="FFFFFF", size=10)
        for col in range(1, 5):
            ws.cell(row[0], col).fill = azul
        row[0] += 1

    def fila(par, unidad, casco, tubos, cab=False):
        for col, val in enumerate([par, unidad, casco, tubos], 1):
            cell = ws.cell(row[0], col, val)
            cell.font = F(bold=cab)
            cell.border = borde
            cell.alignment = Alignment(horizontal="left" if col == 1 else "center", vertical="center",
                                       wrap_text=True)
            if cab:
                cell.fill = gris
        row[0] += 1

    ws.cell(row[0], 1, "HOJA DE ESPECIFICACIONES - INTERCAMBIADOR DE CASCO Y TUBOS (TEMA)").font = F(bold=True, size=13)
    row[0] += 2
    titulo("1. DATOS DE IDENTIFICACIÓN DEL EQUIPO")
    fila("Cliente / Cátedra", "", idn["cliente"], "")
    fila("Ítem / Servicio", "", idn["item"], idn["servicio"])
    fila("Tipo TEMA / Orientación", "", idn["tipo_TEMA"], idn["orientacion"])
    fila("Fecha de emisión", "", str(date.today()), "")
    fila("N° de carcasas en serie / Superficie por carcasa", "- / m2", r["Ns"], f"{r['A_disp']:.2f}")
    fila("Superficie total instalada", "m2", f"{r['A_total']:.2f}", "")
    row[0] += 1
    titulo("2. CONDICIONES DE OPERACIÓN Y PROCESO")
    fila("Parámetro", "Unidad", "LADO CASCO", "LADO TUBOS", cab=True)
    pw = r["propiedades_agua"]
    fila("Fluido", "", h["fluido"], f["fluido"])
    fila("Caudal total", "kg/h", f"{h['W']*3600:,.0f}", f"{bal['Wc']*3600:,.0f}")
    fila("Caudal de vapor entrada / condensado", "kg/h", f"{h['W']*3600:,.0f} / {h['W']*3600:,.0f}", "-")
    fila("Temperatura entrada", "°C", h["T1"], bal["t1"])
    fila("Temperatura salida", "°C", h["T2"], round(bal["t2"], 1))
    fila("Temperatura de saturación", "°C", h["Tsat"], "-")
    fila("Presión de operación / diseño", "kPa abs", f"{h['P_oper']/1e3:.0f} / {h.get('P_diseno', 0)/1e3:.0f}",
         f"{f['P_oper']/1e3:.0f} / {f.get('P_diseno', 0)/1e3:.0f}")
    fila("Temperatura de diseño", "°C", h.get("T_diseno", "-"), f.get("T_diseno", "-"))
    fila("Calor latente", "kJ/kg", f"{h['lam']/1e3:.1f}", "-")
    fila("Densidad vapor / líquido", "kg/m3", f"{h['rho_v']:.3f} / {h['rho_l']:.1f}", f"{pw['rho']:.1f}")
    fila("Viscosidad vapor / líquido", "mPa·s", f"{h['mu_v']*1e3:.4f} / {h['mu_l']*1e3:.4f}", f"{pw['mu']*1e3:.4f}")
    fila("Conductividad térmica vapor / líquido", "W/m·K", f"{h['k_v']:.4f} / {h['k_l']:.3f}", f"{pw['k']:.3f}")
    fila("Calor específico vapor / líquido", "kJ/kg·K", f"{h['cp_v']/1e3:.3f} / {h['cp_l']/1e3:.3f}", f"{pw['cp']/1e3:.3f}")
    fila("Velocidad", "m/s", f"{r['u_vapor_casco']:.1f} (vapor, entre deflectores)", f"{r['v_tubos']:.2f}")
    fila("Resistencia de ensuciamiento requerida", "m2K/W", f"{h['Rd']:.6f}", f"{f['Rd']:.6f}")
    fila("Pérdida de carga admisible", "kPa", f"{h['dP_adm']/1e3:.1f}", f"{f['dP_adm']/1e3:.1f}")
    fila("Pérdida de carga calculada", "kPa", f"{r['dP_casco']/1e3:.2f}", f"{r['dP_tubos']/1e3:.2f}")
    row[0] += 1
    titulo("3. DATOS GEOMÉTRICOS Y CONSTRUCTIVOS DEL HAZ TUBULAR Y CARCASA")
    fila("Parámetro", "Unidad", "CARCASA", "HAZ TUBULAR", cab=True)
    bw = bwg_de(k["Do"], k["Di"])
    fila("Diámetro interno carcasa / Diámetro externo tubo", "mm", f"{r['Ds']*1000:.1f}", f"{k['Do']*1000:.2f}")
    fila("Diámetro interno tubo / BWG", "mm / -", "-", f"{k['Di']*1000:.2f} / BWG {bw}")
    fila("Longitud de tubo", "m", "-", f"{r['L']:.2f}")
    fila("Número de tubos por carcasa", "-", "-", r["Nt"])
    fila("Pasos por tubo", "-", "1", r["n"])
    fila("Arreglo y paso (pitch)", "mm", "-", f"{k['arreglo'].capitalize()} "
         f"{'60°' if k['arreglo']=='triangular' else '90°'}  Pt = {k['Pt']*1000:.2f}")
    fila("Material", "", k["material_casco"], k["material_tubo"])
    fila("Deflectores: tipo / corte", "-", f"Segmentales / {k['corte_deflector']*100:.0f}%", "-")
    fila("Espaciado de deflectores / N° de deflectores", "mm / -", f"{r['B']*1000:.0f} / {r['n_deflectores']}", "-")
    fila("Boquilla entrada vapor / salida condensado (casco)", "DN", r["boq_vapor"][0], r["boq_condensado"][0])
    fila("Boquillas entrada / salida agua (tubos)", "DN", "-", r["boq_tubos"][0])
    row[0] += 1
    titulo("4. RESULTADOS DE DESEMPEÑO TÉRMICO E HIDRÁULICO")
    fila("Parámetro", "Unidad", "VALOR", "OBSERVACIÓN", cab=True)
    fila("Carga térmica Q", "kW", f"{bal['Q']/1e3:,.1f}", "")
    fila("DMLT balanceada x Ft", "K", f"{r['dT_bal']:.2f}", "")
    fila("Factor de corrección Ft", "-", f"{r['Ft']:.3f}", "Ft >= 0.75 " + ("OK" if r["Ft"] >= 0.75 else "NO CUMPLE"))
    fila("ho (casco, condensación)", "W/m2K", f"{r['ho']:,.0f}", r["modelo_ho"])
    fila("hi / hio (tubos, referido al DE)", "W/m2K", f"{r['hi']:,.0f} / {r['hio']:,.0f}", "")
    fila("Uc (limpio)", "W/m2K", f"{r['Uc']:,.0f}", "")
    fila("Ud (diseño)", "W/m2K", f"{r['Ud']:,.0f}", "")
    fila("Rd calculado / Rd requerido", "m2K/W", f"{r['Rd_calc']:.6f} / {r['Rd_req']:.6f}",
         "OK" if r["Rd_calc"] >= r["Rd_req"] else "NO CUMPLE")
    fila("Área requerida limpia / con ensuciamiento / disponible", "m2", f"{r['A_limpia']:.2f} / {r['A_disp']/(1+r['exceso']/100):.2f} / {r['A_disp']:.2f}", "")
    fila("Exceso de área", "%", f"{r['exceso']:.1f}", "")
    fila("Δp tubos: fricción / cabezales / boquillas / total", "kPa",
         f"{r['dP_tubos_fric']/1e3:.2f} / {r['dP_tubos_ret']/1e3:.2f} / {r['dP_tubos_noz']/1e3:.2f} / "
         f"{r['dP_tubos']/1e3:.2f}", "OK" if r["dP_tubos"] <= f["dP_adm"] else "NO CUMPLE")
    fila("Δp casco: haz / boquilla / total", "kPa",
         f"{r['dP_casco_haz']/1e3:.2f} / {r['dP_casco_noz']/1e3:.2f} / {r['dP_casco']/1e3:.2f}",
         "OK" if r["dP_casco"] <= h["dP_adm"] else "NO CUMPLE")
    fila("T pared exterior estimada / T película", "°C", f"{r['Tw']:.1f} / {r['T_film']:.1f}", "")
    fila("VERIFICACIÓN GLOBAL", "", "CUMPLE" if r["cumple"] else "NO CUMPLE", "", cab=True)
    row[0] += 1
    titulo("5. OBSERVACIONES / CRITERIO DE INGENIERÍA")
    for t in notas:
        ws.cell(row[0], 1, "• " + t).font = F()
        ws.merge_cells(start_row=row[0], start_column=1, end_row=row[0], end_column=4)
        ws.cell(row[0], 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row[0]].height = 14 * (1 + len(t) // 95)
        row[0] += 1
    for col, w in zip("ABCD", (52, 12, 34, 34)):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1

    ws2 = wb.create_sheet("Iteraciones del agente")
    ws2.cell(1, 1, "Registro de decisiones del agente").font = Font(name="Arial", bold=True, size=11)
    for i, t in enumerate(log, 3):
        ws2.cell(i, 1, t).font = Font(name="Arial", size=9)
    ws2.column_dimensions["A"].width = 130
    wb.save(ruta)


# =============================================================================
# PROGRAMA PRINCIPAL
# =============================================================================
def main():
    ap = argparse.ArgumentParser(description="Agente de diseño de condensadores de casco y tubos")
    ap.add_argument("caso", nargs="?", help="archivo .json con los inputs (opcional)")
    ap.add_argument("-o", "--salida", default=".", help="carpeta de salida")
    ap.add_argument("--plantilla", action="store_true", help="escribe caso_plantilla.json y termina")
    args = ap.parse_args()
    os.makedirs(args.salida, exist_ok=True)
    if args.plantilla:
        with open(os.path.join(args.salida, "caso_plantilla.json"), "w", encoding="utf-8") as fh:
            json.dump(CASO_EJEMPLO, fh, indent=2, ensure_ascii=False)
        print("Plantilla escrita: caso_plantilla.json")
        return
    if args.caso:
        with open(args.caso, encoding="utf-8") as fh:
            caso = json.load(fh)
    else:
        caso = copy.deepcopy(CASO_EJEMPLO)
    log = []
    try:
        bal, r = disenar(caso, log)
    except ValueError as e:
        print("ERROR en los datos de entrada:", e)
        return
    notas = notas_criticas(caso, r)
    xlsx = os.path.join(args.salida, "Hoja_TEMA_condensador.xlsx")
    escribir_hoja_tema(caso, bal, r, notas, log, xlsx)
    with open(os.path.join(args.salida, "log_agente.txt"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(log))
    salida = {k: v for k, v in r.items() if k not in ("zonas", "propiedades_agua")}
    with open(os.path.join(args.salida, "resultados.json"), "w", encoding="utf-8") as fh:
        json.dump(salida, fh, indent=2, ensure_ascii=False, default=str)
    print("\n".join(log))
    print("\n=== RESULTADO ===")
    print(f"Carcasas: {r['Ns']} | Ds = {r['Ds']*1000:.0f} mm | Nt = {r['Nt']} | pasos = {r['n']} | "
          f"L = {r['L']:.2f} m | B = {r['B']*1000:.0f} mm")
    print(f"ho = {r['ho']:,.0f} | hi = {r['hi']:,.0f} | Uc = {r['Uc']:,.0f} | Ud = {r['Ud']:,.0f} W/m2K")
    print(f"Rd_calc = {r['Rd_calc']:.6f} (req {r['Rd_req']:.6f}) | Ft = {r['Ft']:.3f} | "
          f"exceso de área = {r['exceso']:.1f}%")
    print(f"Δp casco = {r['dP_casco']/1e3:.2f} kPa (adm {caso['caliente']['dP_adm']/1e3:.1f}) | "
          f"Δp tubos = {r['dP_tubos']/1e3:.2f} kPa (adm {caso['fria']['dP_adm']/1e3:.1f}) | "
          f"v tubos = {r['v_tubos']:.2f} m/s")
    print("VERIFICACIÓN:", "CUMPLE" if r["cumple"] else "NO CUMPLE")
    print(f"\nArchivos escritos en: {os.path.abspath(args.salida)}")


if __name__ == "__main__":
    main()
